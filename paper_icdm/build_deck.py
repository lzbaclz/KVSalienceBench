#!/usr/bin/env python3
"""Build the English D2AI deck from frozen v2 evidence and paper PDF figures.

Run from any directory. No experiments are executed and no result JSON is written.
Architecture artwork is rasterized from the draw.io-exported PDF, never redrawn.
"""
from pathlib import Path
import hashlib
import json

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper_icdm"
OUT = PAPER / "output"
ASSETS = OUT / "assets" / "english"
ASSETS.mkdir(parents=True, exist_ok=True)
FONT = "Liberation Sans"
INK, MUTED, BLUE, RED, LIGHT = "202428", "535F6B", "315F82", "8C1D1D", "EDF2F6"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
manifest = []


def load(rel):
    return json.loads((ROOT / "experiments/results" / rel).read_text())


def color(s):
    return RGBColor.from_string(s)


def text(slide, content, x, y, w, h, size=24, bold=False, ink=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.025)
    tf.margin_top = tf.margin_bottom = Inches(0.025)
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.space_before = Pt(0)
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.bold, r.font.color.rgb = bold, color(ink)
    return shape


def rectangle(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color(fill)
    s.line.fill.background()
    return s


def slide(title, subtitle, source, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(s, 0, 0, 0.15, 7.5, BLUE)
    text(s, title, 0.55, 0.34, 12.15, 0.7, 30, True)
    text(s, subtitle, 0.58, 1.09, 12.05, 0.58, 17, ink=MUTED)
    text(s, source, 0.58, 6.77, 11.8, 0.35, 11, ink=MUTED)
    text(s, "D2AI @ IEEE ICDM 2026  |  DM2158", 0.58, 7.15, 11.5, 0.23, 10, ink=MUTED)
    text(s, str(len(prs.slides)), 12.2, 7.12, 0.6, 0.3, 11, ink=MUTED)
    s.notes_slide.notes_text_frame.text = notes
    return s


def callout(s, content, y=5.82, size=22):
    rectangle(s, 0.6, y, 12.1, 0.78, LIGHT)
    text(s, content, 0.8, y + 0.12, 11.7, 0.57, size, True, BLUE)


def table(s, headers, rows, widths, y=2.0, row_h=0.66, size=21):
    shape = s.shapes.add_table(len(rows)+1, len(headers), Inches(0.7), Inches(y),
                               Inches(sum(widths)), Inches(row_h*(len(rows)+1)))
    tab = shape.table
    for col, width in zip(tab.columns, widths):
        col.width = Inches(width)
    for ri, row in enumerate([headers]+rows):
        tab.rows[ri].height = Inches(row_h)
        for ci, value in enumerate(row):
            c = tab.cell(ri, ci)
            c.text = str(value)
            c.margin_left = c.margin_right = Inches(0.12)
            c.margin_top = c.margin_bottom = Inches(0.06)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = color(BLUE if ri == 0 else (LIGHT if ri % 2 else "FFFFFF"))
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.font.bold = ri == 0
                    r.font.color.rgb = color("FFFFFF" if ri == 0 else INK)


def figure(s, name, x, y, w, h):
    pdf = PAPER / "figures" / name
    doc = pymupdf.open(pdf)
    png = ASSETS / (pdf.stem + ".png")
    doc[0].get_pixmap(matrix=pymupdf.Matrix(4, 4), alpha=False).save(png)
    rect = doc[0].rect
    scale = min(w/rect.width, h/rect.height)
    fw, fh = rect.width*scale, rect.height*scale
    s.shapes.add_picture(str(png), Inches(x+(w-fw)/2), Inches(y+(h-fh)/2), Inches(fw), Inches(fh))
    manifest.append((len(prs.slides), name, hashlib.sha256(pdf.read_bytes()).hexdigest()))


v2 = load("icdm_v2.json")
rank = {r["method"]: r for r in v2["pooled_request_split"]["table"]}
quality = load("tost/expand_v2_sensitivity.json")
contrast = quality["contrasts"]["xqp_vs_h2o::pooled::f1_longbench_all_refs"]
cell = contrast["architecture_dataset_cells"]["0.02"]
dataset = contrast["dataset_clusters"]["0.02"]
means = quality["policy_means"]
dec = load("icdm_v2_decomposition.json")

s = slide("Mining Attention Dynamics", "Auditing KV-Saliency Prediction in LLMs", "Camera-ready companion · version-2 main evidence",
          "Introduce this as a measurement audit and benchmark. We study whether a predictor's ranking metric matches the selector's budget and answer-quality objective. No production serving speedup is claimed.")
text(s, "Evaluate the prediction under\nthe decision that will use it.", 0.85, 2.0, 11.7, 1.9, 35, True, BLUE)
text(s, "Ziqing Li · Jiawei Guo · Kaicheng Tang · Jianxi Chen", 0.87, 4.2, 11.5, 0.6, 22)
text(s, "HUST · Beihang University · The Chinese University of Hong Kong", 0.87, 4.93, 11.6, 0.55, 18, ink=MUTED)
text(s, "Corresponding author: Jianxi Chen · chenjx@hust.edu.cn", 0.87, 5.65, 11.4, 0.5, 18, ink=MUTED)

s = slide("Three paths, three evaluation objects", "Feature realization and selector semantics are explicit in the protocol.", "Paper §II and Table I · editable source: evaluation_paths.drawio",
          "The figure separates full-cache offline prediction, a pinned masking loop, and a physical reference. Rows do not imply identical prompt cohorts or execution semantics. The physical reference has different mandatory blocks, budget rules and re-admission behavior. The figure was authored as editable draw.io cells and exported by draw.io.")
figure(s, "evaluation_paths.pdf", 1.0, 1.83, 11.3, 3.8)
callout(s, "Logical retention and physical KV storage require separate measurements.", size=21)

s = slide("Versioned data supports the offline comparison", "Main offline corpus: Llama-3.1-8B-Instruct and Qwen2.5-7B-Instruct", "Source: icdm_v2.json · paper §V-A",
          "The same 128 LongBench QA prompts are collected for each model. Corrected collection aligns queries after RoPE, consumes the prefill-predicted token and generates real lookahead. There are 256 model-requests, not 256 independent source prompts. The source-prompt split holds out the same 32 prompts in both models. Raw traces and model weights are obtained separately.")
table(s, ["Protocol component", "Version 2"], [
    ["Labeled rows", f"{v2['pooled_summary']['n_rows']:,}"],
    ["Source prompts / model-requests", "128 / 256"],
    ["Training / held-out sampled rows", "120,000 / 150,000"],
    ["Additional split control", "Source-prompt disjoint"],
    ["Collector corrections", "RoPE alignment; real lookahead"],
], [5.5, 6.35], y=1.92, row_h=0.64, size=21)
callout(s, "Request and source-prompt units matter more than the row count.")

methods = [("Within-layer EMA", "H2O/attn-EMA"), ("Two-view logistic", "within+cross(2)"), ("LightGBM, balanced", "GBDT(LightGBM)")]
s = slide("Pooled AUC favors the learned predictors", "Two features and one intercept give a compact offline scorer.", "Source: icdm_v2.json / pooled_request_split.table · paper Table II",
          "The refit logistic scorer has three parameters. LightGBM's 150 is a tree count. These are offline fixed-trace comparisons, not a measure of runtime F1. The logistic scorer is close to, but measurably below, LightGBM; this is not an equivalence test.")
table(s, ["Offline score", "Pooled AUC", "Pooled AP"], [[label, f"{rank[k]['auc']:.3f}", f"{rank[k]['auprc']:.3f}"] for label,k in methods], [6.2,2.8,2.85], y=2.17, row_h=0.82, size=24)
callout(s, "A pooled ranking does not enforce a separate budget at every decision.", size=21)

dec_names = {"Within-layer EMA": "within-layer EMA", "Two-view logistic": "two-view logistic (refit)", "LightGBM, balanced": "LightGBM, balanced"}
s = slide("The ordering changes under a decision budget", "Pooled AUC is almost entirely cross-decision pairs; a selector uses the order inside each decision", "Source: icdm_v2_decomposition.json (every held-out row) · paper Table II, lower block",
          "Same-decision pairs are 0.0016 percent of all positive-negative pairs, so pooled AUC is a cross-decision statistic. The binary cross-layer term raises pooled AUC by 0.017 and lowers within-decision AUC by 0.0025 and top-decile recall by 0.046; at 20 percent retention the recall gap shrinks. At the boundary the term promotes previous-layer-hot blocks that are positive 20 percent of the time, against 41 percent for the blocks they displace. This is a measurement of where the reversal comes from in the logistic scorer, not a causal claim about the downstream study.")
table(s, ["Offline score", "Pooled AUC", "Within-decision AUC", "Top-decile recall"],
      [[label, f"{dec['scorers'][dec_names[label]]['auc_pooled']:.3f}", f"{dec['scorers'][dec_names[label]]['auc_same']:.3f}",
        f"{dec['scorers'][dec_names[label]]['recall_0.10']:.3f}"] for label, _ in methods],
      [4.5, 2.3, 2.95, 2.1], y=2.17, row_h=0.82, size=22)
callout(s, "A feature can earn its place pooled and cost recall inside the decision.", size=21)

s = slide("Offset sweep and calibration", "(a) Scaling the cross-layer offset trades pooled AUC for recall; (b) calibration depends on the fit.", "Source: icdm_v2_decomposition.json, icdm_v2.json · Python-generated paper figure",
          "Panel (a) scales the fitted cross-layer offset from zero (the raw within-layer EMA) to twice its value on every held-out row: one eighth of the offset already gives 83 percent of the pooled AUC gain and 65 percent of the top-decile recall loss. Panel (b): the two-view logistic fit has ECE 0.006; unweighted and isotonic-recalibrated LightGBM have 0.005 and 0.003, so the class-balanced tree's poor calibration does not establish a family-level disadvantage. Relevance estimates, in nats, are in the paper text; a query proxy's relevance does not certify query redundancy.")
figure(s, "fig_redund_calib.pdf", 0.9, 1.92, 11.5, 3.82)
callout(s, "Low ECE does not establish a superior retention policy.")

s = slide("The runtime scorer reconstructs its features", "A matched-row bridge separates feature realization from trajectory changes.", "Source: feature_bridge/bridge.{llama.bf16,qwen.fp32}.json · paper §IV-B",
          "All rows use the frozen two-view checkpoint. A and B score the same full-cache rows with different feature realizations. C uses runtime features along its own policy trajectory, but its candidates are only the blocks retained now and at t+h, and its label denominator is taken among those survivors: a survivor-conditioned arm. An evicted block can never be missed there, so the 0.989 is not evidence that retention makes prediction easier, it does not isolate feedback, and it does not measure the coverage lost to eviction. These request-mean AUCs differ in corpus, aggregation and numerics from the historical 0.948-to-0.643 comparison.")
bridges = [load("feature_bridge/bridge.llama.bf16.json")["summary"], load("feature_bridge/bridge.qwen.fp32.json")["summary"]]
table(s, ["Feature / trajectory arm", "Llama bf16", "Qwen fp32"],
      [[label]+[f"{b[key]['mean_request_auc']:.3f}" for b in bridges] for label,key in [
          ("A: offline / full cache", "A_offline_full"), ("B: runtime / full cache", "B_runtime_full"), ("C: runtime / on policy", "C_runtime_policy")]], [6.1,2.85,2.9], y=2.12, row_h=0.76, size=22)
callout(s, "The bridge measures reconstruction; the historical gap's cause remains open.", size=20)

s = slide("Closed-loop answer quality: the pinned rerun", "448 source prompts × 2 models = 896 model–prompt evaluations; 7 datasets", "Source: expand_v2/ provenance and tost/expand_v2_sensitivity.json · paper Exp#8",
          "This is the runtime reconstruction, not the offline predictor. The two architectures answer the same source questions, so the 14 cells are not independent: the primary analysis averages the two models within each dataset and tests over the seven dataset means; the 14 cells are a sensitivity analysis. The two-point margin is a practical tolerance carried over from the archived analysis. It is not pre-registered and not a deployment threshold. The masked loop allows re-admission and does not release physical KV storage.")
text(s, "20% logical retention; each policy drives its own trajectory.\nPinned simulator source and bfloat16 numerics.\nOfficial LongBench scoring with all reference answers.\nPrimary: TOST over seven dataset means; practical ±0.02 F1 margin, not pre-registered.", 0.9, 2.1, 11.55, 3.4, 25)
callout(s, "The main task-quality comparison has two architectures and seven tasks.", size=21)

s = slide("Mean F1 is equivalent within the chosen margin", "Both compressed policies remain below full-cache quality.", "Source: tost/expand_v2_sensitivity.json · paper Table IV",
          "Read the unrounded paired difference, not a difference of rounded table entries. The primary analysis tests over the seven dataset means (p=0.0023 at the two-point margin); the 14-cell analysis, which treats the cells as independent, is the sensitivity analysis. This is mean equivalence between two complete runtime policy realizations in these tasks (they also differ in EMA window and mandatory blocks), not equivalence for every task, no compression loss, a scorer-only intervention, or serving superiority. Allocation comparators are H2O-relative comparisons and not transitive equivalence claims.")
table(s, ["Policy realization", "Mean F1"], [[label,f"{means[key]['f1_longbench_all_refs']['mean']:.3f}"] for label,key in [
    ("Full cache", "full"), ("H2O-style accumulation", "h2o"), ("Reconstructed two-view", "xqp")]], [8.0,3.85], y=1.92, row_h=0.69, size=22)
text(s, f"Paired difference {contrast['grand_mean']:+.4f}   |   seven-dataset TOST p = {dataset['p']:.4f}\nSeven-dataset 90% t interval [{dataset['ci90_t'][0]:+.4f}, {dataset['ci90_t'][1]:+.4f}]\n14-cell sensitivity: p = {cell['p']:.4f}, 90% t interval [{cell['ci90_t'][0]:+.4f}, {cell['ci90_t'][1]:+.4f}]", 0.86, 4.88, 11.7, 1.7, 21, ink=BLUE)

s = slide("Task heterogeneity remains visible", "An average-equivalence result can contain gains and losses by task.", "Source: tost/expand_v2_sensitivity.json · Python-generated paper forest plot",
          "Points are reconstructed-minus-H2O-style F1 differences within model–dataset cells, with 64 paired requests and 95 percent t intervals; the two models of a dataset are adjacent. The filled diamond is the primary seven-dataset estimate with its 90 percent t interval; the hollow diamond is the 14-cell sensitivity estimate. The shaded band is the two-point margin. Do not infer taskwise equivalence from either pooled estimate.")
figure(s, "fig_cell_forest.pdf", 1.8, 1.78, 9.7, 4.75)

s = slide("Physical storage falls; total memory falls much less", "A separate eager, single-request reference backend at 20% block budget", "Source: physical_kv/{llama3.1-v1,qwen25-v1}/summary.json · paper Table III",
          "These values refer to the H2O-style accumulator under matched irreversible eligibility. Llama uses bfloat16 and Qwen uses float32. Full prefill remains uncompressed. Short gates pass, but extended teacher-schedule tests exceed the numerical tolerance. Storage reduction is measured; exact long-run mask/physical equivalence and production throughput are not established.")
phys = [load(f"physical_kv/{n}/summary.json")["cells"] for n in ("llama3.1-v1", "qwen25-v1")]
rows=[]
for label,c in zip(("Llama bf16", "Qwen fp32"),phys):
    a,b=c['h2o_block.masked.b0.2.r0.json'],c['h2o_block.physical.b0.2.r0.json']
    rows.append([label,f"{a['kv_storage_mib_mean']:.0f} → {b['kv_storage_mib_mean']:.0f}",f"{a['peak_decode_gib_mean']:.2f} → {b['peak_decode_gib_mean']:.2f}"])
table(s, ["Model", "Stored KV (MiB)", "Decode peak (GiB)"], rows, [3.75,3.9,4.2], y=2.05, row_h=0.83, size=22)
text(s, "Extended mask/physical parity exceeds the preset numerical tolerance.\nNo production TPOT, concurrency or throughput conclusion follows.", 0.87, 5.02, 11.6, 1.4, 23, ink=RED)

s = slide("Query evidence is scoped to the tested controls", "Phase-aligned dot-max is a probe; it is not Quest's page-bound algorithm.", "Source: physical_kv/*/summary.json / query_v2 · paper §IX",
          "The old collector mixed pre-RoPE queries and post-RoPE cached keys, and its producer was unpinned. We withdraw the faithful-Quest reading. Corrected controls have opposite signs on the two models, so a pooled near-zero should not be read as intrinsic query redundancy. The synthetic needle result is a negative control. The tested cohort cannot support MoE, base-model, 32K-or-longer or sampled-decoding generalization.")
query=[load(f"physical_kv/{n}/summary.json")["query_v2"] for n in ("llama3.1-v1", "qwen25-v1")]
table(s, ["Model", "Request-mean AUC change"], [[n,f"{q['mean_request_auc_delta']:+.4f}"] for n,q in zip(("Llama-3.1-8B", "Qwen2.5-7B"),query)], [6.2,5.65], y=2.14, row_h=0.86, size=24)
callout(s, "No extrapolation to MoE, base models, ≥32K contexts or sampling.", size=21)

s = slide("A benchmark contribution to data systems for AI", "Versioned collection, provenance and explicit evaluation contracts", "Code and paper: github.com/lzbaclz/KVSalienceBench · corresponding author: chenjx@hust.edu.cn",
          "Close with the reproducible measurement contribution. The artifact contains code, immutable experiment records, a pinned MIT-licensed simulator bundle, tests, editable draw.io source and Python chart sources. Model weights, raw traces and source datasets are obtained separately. The default benchmark entry (protocol 2.0) scores pooled and per-decision metrics from a prediction table, without traces or model weights; the shipped example is synthetic. Point readers to the repository and DOI recorded in CITATION.cff once the release is published. Do not call a pending DOI a completed archive.")
text(s, "1. Evaluate scores under the policy's decision budget.\n2. Specify feature reconstruction and selector semantics.\n3. Separate answer quality, physical storage and serving cost.\n4. Preserve provenance and report the limits of each experiment.", 0.87, 2.1, 11.6, 3.6, 25)
callout(s, "The default protocol scores pooled and per-decision metrics together.", size=21)

out = OUT / "final_presentation_en.pptx"
prs.save(out)
reopened = Presentation(out)
assert len(reopened.slides) == 13
for i,s in enumerate(reopened.slides,1):
    assert s.notes_slide.notes_text_frame.text.strip(), f"missing notes on {i}"
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left+shape.width <= prs.slide_width+100
        assert shape.top+shape.height <= prs.slide_height+100
lines = ["# English presentation asset manifest", "", "All architecture artwork comes from the draw.io PDF. Charts come from Python.", "", "| Slide | Paper PDF source | SHA256 |", "|---|---|---|"]
lines += [f"| {s} | `figures/{f}` | `{h}` |" for s,f,h in manifest]
(OUT / "asset_manifest.md").write_text("\n".join(lines)+"\n")
print(f"WROTE {out}: {len(prs.slides)} slides; speaker notes and shape bounds checked")
